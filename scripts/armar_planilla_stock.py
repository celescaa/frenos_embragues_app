"""
Arma la planilla única de stock por proveedor, lista para que alguien del
negocio complete a mano la cantidad que hay de cada producto.

Junta dos cosas que hoy están separadas:

  1. `Planilla_Stock_Por_Proveedor_completa_1.xlsx`, que ya trae ~184.000
     filas de 11 proveedores con descripción, marca y costo, pero
     clasificadas con la taxonomía VIEJA (v2). Subirla tal cual mandaría
     ~110.000 productos al cajón "Varios", porque el importador manda ahí
     todo rubro que no reconoce -- y los rubros v2 (`Suspensión y Dirección`,
     `Correas`, `Transmisión`, `Rodamientos y Mazas`, `Filtros`, `Líquidos`)
     ya no existen. Acá se reclasifican a la taxonomía v3 con
     `clasificar_repuestos.reclasificar_v2()`.
  2. Las listas de proveedores que esa planilla todavía no tiene, leídas de
     sus Excel crudos (cada uno con su propio formato, vía un "adaptador"
     por proveedor, mismo patrón que `limpiar_lista_proveedor.py`) y
     clasificadas de cero con `clasificar_repuestos.clasificar()`.

Lo que NO hace, a propósito:

  - No toca la base de datos. Genera un Excel; cargarlo es el paso aparte
    (`cargar_stock_por_proveedor.py`), que además solo carga las filas con
    cantidad completada.
  - No inventa códigos de barras. Ninguna de las listas de proveedores trae
    uno (se revisaron las 26, columna por columna): los códigos numéricos
    que parecen EAN son códigos de fábrica. La columna queda vacía para
    completar escaneando la caja.
  - No completa la cantidad en stock. Ese es justamente el dato que el
    negocio tiene que poner a mano y ninguna lista de precios contiene.

Uso (desde la raíz del proyecto):
    python scripts/armar_planilla_stock.py [--listas CARPETA] [--salida ARCHIVO]
"""
import argparse
import os
import sys
from decimal import Decimal, ROUND_HALF_UP

from openpyxl import Workbook, load_workbook
from openpyxl.cell import WriteOnlyCell
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from scripts import clasificar_repuestos as clasif
from scripts import nomenclatura as nom
from scripts.importar_datos import _limpiar, _numero

CARPETA_LISTAS_POR_DEFECTO = os.path.expanduser(
    "~/Downloads/Lista De Precios San Ignacio"
)
SALIDA_POR_DEFECTO = "plantillas/Planilla_Stock_Completa.xlsx"

# El margen con el que se sugiere el precio de venta sobre el costo. Es el
# mismo 30% que ya usan `limpiar_lista_proveedor.py` y el alta rápida de
# producto desde una compra, para no tener tres márgenes distintos dando
# vueltas. Queda editable en la planilla.
MARGEN = Decimal("1.30")

# El orden de las columnas NO es decorativo: `cargar_stock_por_proveedor.py`
# lee cada campo por posición (fila[0], fila[1], ...), así que mover una
# columna acá carga los datos en el campo equivocado sin ningún error.
COLUMNAS = [
    ("Cargar (SI/NO)", 14),
    ("Código interno", 22),
    ("Categoría", 20),
    ("Descripción / Nombre (*)", 52),
    ("Marca", 18),
    ("Modelo compatible", 34),
    ("Precio costo (*)", 15),
    ("Precio venta (*)", 15),
    ("Cantidad en stock (*)", 20),
    ("Stock mínimo", 13),
    ("Código de barras", 18),
    ("Subcategoría", 30),
]

STOCK_MINIMO_POR_DEFECTO = 2

RELLENO_ENCABEZADO = PatternFill("solid", fgColor="1F1F1F")
RELLENO_A_COMPLETAR = PatternFill("solid", fgColor="F2B705")
FUENTE_ENCABEZADO = Font(bold=True, color="FFFFFF", size=11)
FUENTE_A_COMPLETAR = Font(bold=True, color="1F1F1F", size=11)

INSTRUCCIONES = [
    ("Planilla de stock por proveedor — Repuestos San Ignacio", True),
    ("", False),
    ("Cada hoja es un proveedor. Adentro está TODO el catálogo que ese", False),
    ("proveedor vende, no solo lo que hay en el local.", False),
    ("", False),
    ("QUÉ HAY QUE COMPLETAR", True),
    ("", False),
    ("Una sola columna es obligatoria: 'Cantidad en stock'.", False),
    ("Poné la cantidad SOLO en los productos que el local realmente tiene.", False),
    ("Las filas que queden sin cantidad NO se cargan al sistema: se toman", False),
    ("como productos que el proveedor vende pero el local no trabaja.", False),
    ("Por eso no hace falta recorrer las miles de filas de cada hoja: se", False),
    ("busca con Ctrl+F (o Cmd+F) el código o la descripción de lo que hay", False),
    ("en el estante y se le pone el número al lado.", False),
    ("", False),
    ("LO DEMÁS ES OPCIONAL", True),
    ("", False),
    ("Precio de venta: viene sugerido con 30% sobre el costo. Si el margen", False),
    ("de ese producto es otro, se pisa el número y listo.", False),
    ("Código de barras: viene vacío. Ninguna lista de proveedor lo trae —", False),
    ("los códigos largos que parecen código de barras son códigos de", False),
    ("fábrica. Se puede dejar vacío y escanear la caja después desde la", False),
    ("ficha del producto, que es más rápido y no se presta a error de tipeo.", False),
    ("Stock mínimo: viene en 2. Es a partir de cuánto el producto aparece", False),
    ("en la lista de pedidos. Poné 0 para no controlarle la reposición.", False),
    ("Categoría y subcategoría: vienen completas. Corregí las que veas mal.", False),
    ("", False),
    ("CARGAR (SI/NO)", True),
    ("", False),
    ("Viene en SI. Poné NO en una fila para que no se cargue aunque tenga", False),
    ("cantidad — por ejemplo si está duplicada con otro proveedor.", False),
    ("", False),
    ("CÓMO SE SUBE AL SISTEMA", True),
    ("", False),
    ("Desde la raíz del proyecto:", False),
    ("    python scripts/cargar_stock_por_proveedor.py <este archivo>", False),
    ("", False),
    ("Se puede correr más de una vez: si el producto ya existe lo actualiza", False),
    ("en vez de duplicarlo. La cantidad REEMPLAZA el stock del sistema (no", False),
    ("lo suma), porque esto carga el estado real del estante.", False),
]


# --- Lectura de los Excel crudos --------------------------------------------

def _filas_xlsx(ruta, hoja=0, desde=1):
    """Filas de una hoja de un .xlsx, salteando `desde` filas de encabezado."""
    wb = load_workbook(ruta, read_only=True, data_only=True)
    ws = wb.worksheets[hoja] if isinstance(hoja, int) else wb[hoja]
    for i, fila in enumerate(ws.iter_rows(values_only=True)):
        if i < desde:
            continue
        yield fila
    wb.close()


def _filas_xls(ruta, hoja=0, desde=1):
    """Ídem para los .xls viejos (formato binario, openpyxl no los abre)."""
    import xlrd
    wb = xlrd.open_workbook(ruta)
    ws = wb.sheet_by_index(hoja)
    for i in range(desde, ws.nrows):
        yield tuple(ws.row_values(i))


def _codigo(valor):
    """Normaliza un código que Excel pudo haber leído como número.

    Sin esto, el código 1199 de Comfrix llega como el float 1199.0 y se
    guardaría literalmente como "1199.0", que no matchea nada cuando alguien
    lo busca.
    """
    texto = _limpiar(valor)
    if texto.endswith(".0") and texto[:-2].isdigit():
        return texto[:-2]
    return texto


# --- Adaptadores: un formato de proveedor -> filas de la planilla ------------
# Cada adaptador recibe la ruta del Excel crudo y devuelve dicts con las
# claves de la planilla. La clasificación se hace acá porque cada proveedor
# tiene su propia columna útil: Michelli trae el tipo de pieza en 'Producto',
# Devoto lo trae en 'UBICACION', y ninguno de los dos lo dice en la
# descripción sola.

def adaptador_michelli(ruta):
    for fila in _filas_xlsx(ruta, desde=2):
        codigo_fabrica, codigo_interno, fabricante = fila[0], fila[1], fila[2]
        producto, descripcion = fila[3], fila[4]
        # 'Costo Neto c/IVA': mismo criterio que ya se usó con Distrisuper y
        # Eine, donde el costo comparable es el que incluye IVA.
        costo = _numero(fila[7])
        nombre = " ".join(x for x in (_limpiar(producto), _limpiar(descripcion)) if x)
        if not nombre:
            continue
        categoria, subcategoria = clasif.clasificar(producto, descripcion)
        yield {
            "codigo": _codigo(codigo_fabrica) or _codigo(codigo_interno),
            "nombre": nombre,
            "categoria": categoria,
            "subcategoria": subcategoria,
            "marca": _limpiar(fabricante),
            "modelo": _limpiar(descripcion),
            "costo": costo,
        }


def adaptador_devoto(ruta):
    for fila in _filas_xlsx(ruta, desde=1):
        codigo, marca_prod, marca_veh = fila[0], fila[1], fila[2]
        vehiculo, ubicacion, descripcion = fila[3], fila[4], fila[5]
        nombre = _limpiar(descripcion)
        if not nombre:
            continue
        categoria, subcategoria = clasif.clasificar(ubicacion, descripcion)
        modelo = " ".join(x for x in (_limpiar(marca_veh), _limpiar(vehiculo)) if x)
        yield {
            "codigo": _codigo(codigo),
            "nombre": nombre,
            "categoria": categoria,
            "subcategoria": subcategoria,
            "marca": _limpiar(marca_prod),
            "modelo": modelo,
            "costo": _numero(fila[7]),
        }


def adaptador_zapatas_sensei(ruta):
    """Zapatas de freno de Sen-Sei: DOS productos por fila.

    La misma fila lista la zapata LPR (código y precio en las columnas 0 y 1)
    y su equivalente Comfrix (columnas 4 y 5), compartiendo la descripción.
    Son dos productos distintos de dos marcas distintas, así que salen como
    dos filas -- si no, se pierde una de las dos marcas.
    """
    for fila in _filas_xls(ruta, desde=6):
        descripcion = _limpiar(fila[3])
        if not descripcion:
            continue
        categoria, subcategoria = clasif.clasificar("zapata de freno", descripcion)
        for columna_codigo, columna_precio, marca in ((0, 1, "LPR"), (4, 5, "COMFRIX")):
            codigo = _codigo(fila[columna_codigo])
            costo = _numero(fila[columna_precio])
            if not codigo or not costo:
                continue
            yield {
                "codigo": f"{marca} {codigo}" if not codigo.upper().startswith(marca) else codigo,
                "nombre": f"ZAPATA DE FRENO {descripcion}",
                "categoria": categoria,
                "subcategoria": subcategoria,
                "marca": marca,
                "modelo": descripcion,
                "costo": costo,
            }


# Proveedor -> (archivo dentro de la carpeta de listas, adaptador).
# Solo los que la planilla base todavía no tiene. Los otros archivos sueltos
# de la carpeta (los 3 de Papierttai, el de Rodamitre, las pastillas MAG de
# Sen-Sei, roncal, Eine, Distrisuper, fox de Rio, Omar, Zerbini, infofren) ya
# están adentro de la planilla base -- se verificó cruzando sus códigos.
PROVEEDORES_NUEVOS = {
    "Michelli": ("Nuevas/61806 10082026_Michelli.xlsx", adaptador_michelli),
    "Deboto": ("Nuevas/bd 02agosto2026_devoto.xlsx", adaptador_devoto),
}

# Filas que se suman a un proveedor que YA está en la planilla base.
AGREGADOS = {
    "Miguel Angel Sen-Sei": [
        ("Nuevas/LISTA ZAPATAS DE FRENO AGOSTO 2026 (1) sensei.xls", adaptador_zapatas_sensei),
    ],
}


# --- Armado -----------------------------------------------------------------

def precio_venta_sugerido(costo):
    if not costo:
        return None
    return (Decimal(costo) * MARGEN).quantize(Decimal("0.01"), rounding=ROUND_HALF_UP)


def filas_de_planilla_base(ruta, resumen):
    """Lee la planilla ya clasificada con v2 y la reclasifica a v3."""
    wb = load_workbook(ruta, read_only=True, data_only=True)
    for ws in wb.worksheets:
        if ws.title in ("INSTRUCCIONES", "Listas"):
            continue
        productos = []
        # El encabezado se busca, no se asume en la primera fila: la hoja
        # SIN IDENTIFICAR lleva un aviso arriba de los títulos, y dando por
        # sentado que la fila 1 es el encabezado se colaba la fila de títulos
        # como si fuera un producto llamado "Descripción / Nombre (*)". Es el
        # mismo criterio que usa `cargar_stock_por_proveedor._filas_de_hoja`.
        encabezado_visto = False
        for fila in ws.iter_rows(max_col=12, values_only=True):
            if not encabezado_visto:
                encabezado_visto = _limpiar(fila[0]).upper().startswith("CARGAR")
                continue
            nombre = _limpiar(fila[3])
            if not nombre:
                continue
            categoria, subcategoria = clasif.reclasificar_v2(
                fila[2], fila[11] if len(fila) > 11 else None,
                nombre, fila[4], fila[5],
            )
            resumen["reclasificadas"] += 1
            productos.append({
                "codigo": _codigo(fila[1]),
                "nombre": nombre,
                "categoria": categoria,
                "subcategoria": subcategoria,
                "marca": _limpiar(fila[4]),
                "modelo": _limpiar(fila[5]),
                "costo": _numero(fila[6]),
            })
        yield ws.title, productos
    wb.close()


# Tope de caracteres de `modelo_compatible` al consolidar. La lista de Devoto
# llega a repetir un mismo código 58 veces, una por auto; pegar los 58 haría
# una celda ilegible y una columna imposible de mostrar en la tabla de Stock.
LARGO_MAXIMO_MODELO = 200


def consolidar(productos):
    """Junta en un solo producto las filas que son el mismo código repetido.

    Varias listas traen una fila POR AUTO en vez de una por producto (Devoto
    es el caso extremo: 24.320 filas para 9.858 productos). Quedarse con la
    primera y descartar el resto -- que es lo que hay que hacer igual, porque
    `productos.codigo` es UNIQUE -- tiraría los otros 57 autos de cada
    producto, que es justamente el dato con el que después se busca "pastilla
    palio" en el sistema. Así que en vez de descartarlas, sus modelos se
    acumulan en `modelo_compatible`.

    Devuelve (productos consolidados, cuántas filas se fusionaron).
    """
    consolidados = {}
    orden = []
    fusionadas = 0
    for p in productos:
        clave = p["codigo"] or p["nombre"]
        anterior = consolidados.get(clave)
        if anterior is None:
            consolidados[clave] = dict(p, modelos=[p["modelo"]] if p["modelo"] else [])
            orden.append(clave)
            continue
        fusionadas += 1
        if p["modelo"] and p["modelo"] not in anterior["modelos"]:
            anterior["modelos"].append(p["modelo"])
        # El costo puede venir vacío en una de las repeticiones: se conserva
        # el primero que exista en vez de dejar el producto sin precio.
        if not anterior["costo"] and p["costo"]:
            anterior["costo"] = p["costo"]

    salida = []
    for clave in orden:
        p = consolidados[clave]
        modelo = " / ".join(p.pop("modelos"))
        if len(modelo) > LARGO_MAXIMO_MODELO:
            modelo = modelo[:LARGO_MAXIMO_MODELO].rsplit(" / ", 1)[0] + " / y otros"
        p["modelo"] = modelo
        salida.append(p)
    return salida, fusionadas


def escribir_hoja(wb, nombre, productos, resumen):
    ws = wb.create_sheet(title=nombre[:31])
    for i, (titulo, ancho) in enumerate(COLUMNAS, start=1):
        ws.column_dimensions[get_column_letter(i)].width = ancho
    ws.freeze_panes = "A2"

    encabezado = []
    for titulo, _ in COLUMNAS:
        celda = WriteOnlyCell(ws, value=titulo)
        # La única columna obligatoria va en dorado para que se distinga de
        # las otras once de un vistazo: la planilla tiene miles de filas y
        # quien la completa no debería tener que releer las instrucciones
        # para acordarse de cuál es.
        a_completar = titulo.startswith("Cantidad en stock")
        celda.fill = RELLENO_A_COMPLETAR if a_completar else RELLENO_ENCABEZADO
        celda.font = FUENTE_A_COMPLETAR if a_completar else FUENTE_ENCABEZADO
        encabezado.append(celda)
    ws.append(encabezado)

    consolidados, fusionadas = consolidar(productos)
    resumen["fusionadas"] += fusionadas
    for p in consolidados:
        costo = p["costo"] or None
        venta = precio_venta_sugerido(costo)
        if not costo:
            resumen["sin_costo"] += 1
        # La nomenclatura se aplica ya al generar, así la planilla nace
        # prolija y no necesita pasar después por normalizar_planilla_stock.py
        # (que existe para las planillas que ya están en manos del negocio).
        limpio = nom.normalizar_fila(
            proveedor=ws.title, descripcion=p["nombre"], marca=p["marca"],
            rubro=p["categoria"], subrubro=p["subcategoria"],
            modelo=p["modelo"], codigo=p["codigo"] or "",
        )
        # Cuenta el rubro que de verdad se escribe, no el que traía antes de
        # normalizar.
        if limpio["rubro"] == clasif.VARIOS:
            resumen["en_varios"] += 1
        ws.append([
            "SI",
            p["codigo"] or None,
            limpio["rubro"],
            limpio["nombre"],
            limpio["marca"] or None,
            limpio["modelo"] or None,
            float(costo) if costo else None,
            float(venta) if venta else None,
            None,                       # Cantidad en stock: lo completa el negocio
            STOCK_MINIMO_POR_DEFECTO,
            None,                       # Código de barras: ninguna lista lo trae
            limpio["subrubro"],
        ])
        resumen["filas"] += 1
    return ws


def escribir_instrucciones(wb):
    ws = wb.create_sheet(title="INSTRUCCIONES")
    ws.column_dimensions["A"].width = 78
    for texto, es_titulo in INSTRUCCIONES:
        celda = WriteOnlyCell(ws, value=texto)
        if es_titulo:
            celda.font = Font(bold=True, size=12)
        ws.append([celda])
    return ws


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--listas", default=CARPETA_LISTAS_POR_DEFECTO,
                        help="carpeta con los Excel crudos de los proveedores")
    parser.add_argument("--salida", default=SALIDA_POR_DEFECTO)
    args = parser.parse_args()

    base = os.path.join(args.listas, "Planilla_Stock_Por_Proveedor_completa_1.xlsx")
    if not os.path.exists(base):
        parser.error(f"no encontré la planilla base en {base}")

    resumen = {"filas": 0, "reclasificadas": 0, "fusionadas": 0,
               "sin_costo": 0, "en_varios": 0}

    # write_only: con ~220.000 filas, escribir celda por celda no entra en el
    # tiempo ni en la memoria disponibles (ya documentado para
    # extraer_catalogo_referencia.py, mismo volumen).
    wb = Workbook(write_only=True)
    escribir_instrucciones(wb)

    por_proveedor = {}
    print("Leyendo la planilla base y reclasificando a la taxonomía v3...")
    for proveedor, productos in filas_de_planilla_base(base, resumen):
        por_proveedor.setdefault(proveedor, []).extend(productos)

    for proveedor, entradas in AGREGADOS.items():
        for archivo, adaptador in entradas:
            ruta = os.path.join(args.listas, archivo)
            if not os.path.exists(ruta):
                print(f"  aviso: falta {archivo}, salteado")
                continue
            nuevas = list(adaptador(ruta))
            por_proveedor.setdefault(proveedor, []).extend(nuevas)
            print(f"  + {len(nuevas)} filas para {proveedor} desde {os.path.basename(archivo)}")

    for proveedor, (archivo, adaptador) in PROVEEDORES_NUEVOS.items():
        ruta = os.path.join(args.listas, archivo)
        if not os.path.exists(ruta):
            print(f"  aviso: falta {archivo}, {proveedor} queda vacío")
            continue
        nuevas = list(adaptador(ruta))
        por_proveedor.setdefault(proveedor, []).extend(nuevas)
        print(f"  + {len(nuevas)} filas para {proveedor} desde {os.path.basename(archivo)}")

    print("\nEscribiendo hojas:")
    for proveedor in sorted(por_proveedor, key=lambda p: -len(por_proveedor[p])):
        productos = por_proveedor[proveedor]
        antes = resumen["filas"]
        escribir_hoja(wb, proveedor, productos, resumen)
        print(f"  {proveedor:26} {resumen['filas'] - antes:7} filas")

    os.makedirs(os.path.dirname(os.path.abspath(args.salida)) or ".", exist_ok=True)
    wb.save(args.salida)

    print()
    print("=" * 60)
    print("  PLANILLA GENERADA")
    print("=" * 60)
    print(f"  Archivo: {args.salida}")
    print(f"  Filas totales: {resumen['filas']:,}".replace(",", "."))
    print(f"  Reclasificadas de la taxonomía vieja: {resumen['reclasificadas']:,}".replace(",", "."))
    print(f"  Filas del mismo código fusionadas: {resumen['fusionadas']:,}".replace(",", "."))
    print(f"  Sin precio de costo (quedan sin venta sugerida): {resumen['sin_costo']:,}".replace(",", "."))
    print(f"  En el rubro 'Varios': {resumen['en_varios']:,}".replace(",", "."))
    print("=" * 60)
    print("\nFalta completar a mano la columna 'Cantidad en stock'.")
    print("Las filas sin cantidad no se cargan al sistema.")


if __name__ == "__main__":
    main()
