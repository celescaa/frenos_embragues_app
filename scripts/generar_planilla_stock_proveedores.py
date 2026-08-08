"""
Genera una planilla para que alguien del negocio (sin acceso al sistema)
cargue a mano el stock real de productos, organizada por proveedor: una
hoja por cada proveedor ya cargado en el sistema, más una hoja aparte para
anotar productos de un proveedor que todavía no se identificó (por ejemplo,
una lista de precios de la que no se acuerdan a quién pertenece).

Pensada para el mismo caso de uso que limpiar_lista_proveedor.py (una lista
de precios de un proveedor), pero al revés: acá no hay un Excel de origen
para limpiar, se arranca de cero para que alguien complete a mano qué tiene
en stock, con qué cantidad y a qué precio — el resultado ya queda listo
para importarse con cargar_stock_por_proveedor.py.

Uso (desde la raíz del proyecto): python scripts/generar_planilla_stock_proveedores.py [archivo_salida.xlsx]
  Por defecto escribe en plantillas/Planilla_Stock_Por_Proveedor.xlsx.
"""
import sys
import os
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

# database.py vive en el paquete core/, en la raíz del proyecto (un nivel
# arriba de scripts/).
sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from core import database as db

SALIDA_POR_DEFECTO = "plantillas/Planilla_Stock_Por_Proveedor.xlsx"

FILAS_VACIAS = 60  # cuántas filas en blanco dejar listas para completar en cada hoja

HEADERS = [
    "Cargar (SI/NO)", "Código interno", "Categoría", "Descripción / Nombre (*)",
    "Marca", "Modelo compatible", "Precio costo (*)", "Precio venta (*)",
    "Cantidad en stock (*)", "Stock mínimo", "Código de barras",
]
ANCHOS = [12, 16, 12, 45, 16, 20, 14, 14, 16, 12, 16]


def _escribir_hoja(wb, nombre_hoja, categorias, con_nota_identificacion=False):
    nombre_hoja = nombre_hoja[:31] if len(nombre_hoja) <= 31 else nombre_hoja[:28] + "..."
    ws = wb.create_sheet(nombre_hoja)

    fila_actual = 1
    if con_nota_identificacion:
        aviso = (
            "Productos de una lista de precios que todavía no identificamos a qué proveedor pertenece. "
            "Si te acordás o encontrás un dato (CUIT, nombre, teléfono) que la identifique, mejor cargarlo "
            "directamente en la hoja de ese proveedor — esta hoja es para lo que quede sin identificar."
        )
        ws.merge_cells("A1:K1")
        ws["A1"] = aviso
        ws["A1"].font = Font(bold=True, color="9C0006")
        ws["A1"].fill = PatternFill("solid", fgColor="FFC7CE")
        ws["A1"].alignment = Alignment(wrap_text=True, vertical="top")
        ws.row_dimensions[1].height = 45
        fila_actual = 3

    header_row = fila_actual
    for col, h in enumerate(HEADERS, start=1):
        c = ws.cell(row=header_row, column=col, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="111111")

    primera_fila_datos = header_row + 1
    ultima_fila_datos = header_row + FILAS_VACIAS
    for i in range(primera_fila_datos, ultima_fila_datos + 1):
        ws.cell(row=i, column=1, value="SI")
        ws.cell(row=i, column=10, value=2)  # stock mínimo sugerido por defecto

    dv_si_no = DataValidation(type="list", formula1='"SI,NO"', allow_blank=False)
    ws.add_data_validation(dv_si_no)
    dv_si_no.add(f"A{primera_fila_datos}:A{ultima_fila_datos}")

    dv_categoria = DataValidation(
        type="list", formula1=f'"{",".join(categorias)}"', allow_blank=True
    )
    ws.add_data_validation(dv_categoria)
    dv_categoria.add(f"C{primera_fila_datos}:C{ultima_fila_datos}")

    for col, w in enumerate(ANCHOS, start=1):
        ws.column_dimensions[get_column_letter(col)].width = w

    ws.freeze_panes = f"A{primera_fila_datos}"
    ws.auto_filter.ref = f"A{header_row}:K{ultima_fila_datos}"
    return ws


def generar(salida=SALIDA_POR_DEFECTO):
    db.init_db()
    conn = db.get_connection()
    proveedores = conn.execute("SELECT nombre FROM proveedores WHERE activo=1 ORDER BY nombre").fetchall()
    categorias = db.obtener_categorias(conn)
    conn.close()

    wb = openpyxl.Workbook()
    wb.remove(wb.active)

    ws_instr = wb.create_sheet("INSTRUCCIONES", 0)
    ws_instr["A1"] = "Cómo completar esta planilla"
    ws_instr["A1"].font = Font(bold=True, size=14)
    instrucciones = [
        "",
        "Cada hoja es un proveedor de los que ya están cargados en el sistema. Hay una hoja extra,",
        "'SIN IDENTIFICAR', para productos de una lista de precios que no sepamos de qué proveedor es.",
        "",
        "1. Buscá la hoja del proveedor correspondiente (o SIN IDENTIFICAR si no estás segura/o).",
        "2. Cargá una fila por producto: Descripción, Categoría, Precio costo, Precio venta y",
        "   Cantidad en stock son los datos importantes. Código interno, Marca, Modelo compatible,",
        "   Stock mínimo y Código de barras son opcionales (dejalos vacíos si no los tenés).",
        "3. La columna 'Cargar (SI/NO)' ya viene en SI. Si hay algún producto que no querés sumar",
        "   al sistema todavía, poné NO en esa fila (no se va a cargar, pero queda anotado).",
        "4. 'Cantidad en stock' es lo que hay ahora mismo en el local — es lo que va a quedar",
        "   como stock inicial de ese producto en el sistema.",
        "5. 'Stock mínimo' ya viene sugerido en 2 (avisa para reponer cuando quede por debajo). Se",
        "   puede cambiar por producto, o poner 0 si todavía no querés que el sistema lo controle.",
        "6. Guardá el archivo y devolvélo — desde ahí se carga todo junto al sistema.",
    ]
    for i, linea in enumerate(instrucciones, start=2):
        ws_instr.cell(row=i, column=1, value=linea)
    ws_instr.column_dimensions["A"].width = 100

    for p in proveedores:
        _escribir_hoja(wb, p["nombre"], categorias)

    _escribir_hoja(wb, "SIN IDENTIFICAR", categorias, con_nota_identificacion=True)

    wb.save(salida)
    print(f"Guardado en: {salida} ({len(proveedores)} proveedores + hoja SIN IDENTIFICAR)")


if __name__ == "__main__":
    generar(sys.argv[1] if len(sys.argv) > 1 else SALIDA_POR_DEFECTO)
