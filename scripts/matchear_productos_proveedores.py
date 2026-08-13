"""
Puebla automáticamente producto_proveedor (comparador de precios de
proveedores, ver /pedidos) matcheando las listas de precios de varios
proveedores contra el catálogo ya cargado en el sistema (`productos`).

No crea productos nuevos ni toca stock: solo vincula, para un producto que
YA existe en la base, a qué precio se lo vende cada proveedor cuya lista
se procese acá. Es exactamente la tabla que ya usa /pedidos para elegir el
proveedor más barato — con esto se puebla sola en vez de cargarla a mano
producto por producto desde la ficha.

Formato de entrada: un Excel con una hoja por proveedor (incluida la hoja
"SIN IDENTIFICAR" si existe, tratada como el proveedor placeholder que ya
usa cargar_stock_por_proveedor.py) y encabezados en la primera fila. Sirve
tanto la planilla de generar_planilla_stock_proveedores.py (que trae código
de barras) como el catálogo de referencia de extraer_catalogo_referencia.py
(que no lo trae — en ese caso el matching cae directo al fallback por
descripción para todas las filas). Los encabezados se detectan por nombre,
no por posición fija, así que alcanza con que digan lo mismo aunque el
orden de columnas cambie.

Matching, en orden:
  1. Código de barras exacto (si la columna existe y la fila lo trae).
  2. Descripción normalizada (sin acentos, mayúsculas) + misma categoría,
     por similitud de texto (difflib, mismo criterio que ya usa
     importar_factura.py para las líneas de una factura). Umbral configurable.
  Sin match en ninguno de los dos: la fila queda afuera y aparece en el
  reporte para revisión manual.

Antes de correr esto a escala completa: confirmar a mano, columna por
columna, si las planillas fuente realmente traen código de barras real (no
vacío/inventado) — si no lo traen, todo cae al fallback por descripción y
conviene arrancar con --categoria en una categoría chica (ej. "Frenos" con
--subcategoria "Pastillas") como piloto antes de correrlo sobre el archivo
completo, para poder auditar el reporte de salida antes de confiar en el
resultado a escala.

Uso (desde la raíz del proyecto):
  python scripts/matchear_productos_proveedores.py <archivo.xlsx> [--categoria "Frenos"] [--subcategoria "Pastillas"] [--umbral 0.6] [--aplicar]

Sin --aplicar corre en modo simulación (dry-run): solo imprime y guarda el
reporte, no escribe nada en la base. Con --aplicar además inserta/actualiza
producto_proveedor.
"""
import sys
import os
import argparse
import difflib
import unicodedata

import openpyxl

sys.path.insert(0, os.path.dirname(os.path.dirname(os.path.abspath(__file__))))
from core import database as db

REPORTE_POR_DEFECTO = "listas_proveedores/Reporte_Matching_Proveedores.xlsx"
NOMBRE_PROVEEDOR_SIN_IDENTIFICAR = "Proveedor sin identificar (revisar)"
HOJAS_IGNORADAS = {"INSTRUCCIONES", "INFO"}
UMBRAL_TEXTO_DEFAULT = 0.6

# Nombres de encabezado aceptados por columna (case/acentos ya normalizados
# al comparar) — cubre tanto generar_planilla_stock_proveedores.py como
# extraer_catalogo_referencia.py, que no usan exactamente los mismos títulos.
ENCABEZADOS = {
    "codigo": ["codigo interno", "codigo"],
    "categoria": ["categoria"],
    "subcategoria": ["subcategoria"],
    "descripcion": ["descripcion / nombre (*)", "descripcion"],
    "marca": ["marca"],
    "modelo": ["modelo compatible"],
    "costo": ["precio costo (*)", "precio costo"],
    "codigo_barras": ["codigo de barras", "codigo de barras (*)"],
}


def _sin_acentos(s):
    if s is None:
        return ""
    return "".join(c for c in unicodedata.normalize("NFD", str(s)) if unicodedata.category(c) != "Mn").lower().strip()


def _limpiar(v):
    return str(v).strip() if v is not None else ""


def _numero(v):
    if v is None or v == "":
        return 0.0
    if isinstance(v, (int, float)):
        return float(v)
    texto = str(v).strip().replace("$", "").replace(" ", "")
    if "," in texto and "." in texto:
        texto = texto.replace(".", "").replace(",", ".")
    elif "," in texto:
        texto = texto.replace(",", ".")
    try:
        return float(texto)
    except ValueError:
        return 0.0


def _mapear_encabezados(headers):
    idx = {}
    for i, h in enumerate(headers or []):
        h_norm = _sin_acentos(h)
        for campo, alias in ENCABEZADOS.items():
            if h_norm in alias:
                idx[campo] = i
    return idx


def _obtener_o_crear_proveedor(conn, nombre, resumen):
    fila = conn.execute("SELECT id FROM proveedores WHERE nombre = ?", (nombre,)).fetchone()
    if fila:
        return fila["id"]
    resumen["proveedores_nuevos"].append(nombre)
    cur = conn.execute("INSERT INTO proveedores (nombre) VALUES (?)", (nombre,))
    return cur.lastrowid


def _catalogo_candidatos(conn, categoria=None, subcategoria=None):
    """Productos ya cargados en el sistema, con texto normalizado para
    matchear por descripción y un índice por código de barras."""
    consulta = "SELECT * FROM productos"
    condiciones, params = [], []
    if categoria:
        condiciones.append("categoria = ?")
        params.append(categoria)
    if subcategoria:
        condiciones.append("subcategoria = ?")
        params.append(subcategoria)
    if condiciones:
        consulta += " WHERE " + " AND ".join(condiciones)
    filas = conn.execute(consulta, params).fetchall()

    candidatos = []
    por_codigo_barras = {}
    for p in filas:
        texto_norm = _sin_acentos(f"{p['nombre']} {p['marca'] or ''} {p['modelo_compatible'] or ''}")
        candidatos.append({"id": p["id"], "nombre": p["nombre"], "categoria": p["categoria"], "texto_norm": texto_norm})
        if p["codigo_barras"]:
            por_codigo_barras[_limpiar(p["codigo_barras"])] = p["id"]
    return candidatos, por_codigo_barras


def _matchear_por_descripcion(descripcion, categoria, candidatos, umbral):
    desc_norm = _sin_acentos(descripcion)
    mejor, mejor_score = None, 0
    for c in candidatos:
        if categoria and c["categoria"] and _sin_acentos(c["categoria"]) != _sin_acentos(categoria):
            continue
        score = difflib.SequenceMatcher(None, desc_norm, c["texto_norm"]).ratio()
        if score > mejor_score:
            mejor_score, mejor = score, c
    if mejor and mejor_score >= umbral:
        return mejor, mejor_score
    return None, mejor_score


def procesar(archivo, categoria, subcategoria, umbral, aplicar):
    wb = openpyxl.load_workbook(archivo, data_only=True, read_only=True)
    conn = db.get_connection()

    candidatos, por_codigo_barras = _catalogo_candidatos(conn, categoria, subcategoria)
    if not candidatos:
        print("No hay productos cargados en el catálogo (con ese filtro de categoría/subcategoría, si se usó). Nada para matchear.")
        conn.close()
        return

    resumen = {"proveedores_nuevos": [], "por_codigo_barras": 0, "por_descripcion": 0, "sin_match": 0}
    filas_reporte = []

    for nombre_hoja in wb.sheetnames:
        if nombre_hoja in HOJAS_IGNORADAS:
            continue
        ws = wb[nombre_hoja]
        filas_iter = ws.iter_rows(values_only=True)
        headers = next(filas_iter, None)
        idx = _mapear_encabezados(headers)
        if "descripcion" not in idx:
            continue  # esta hoja no tiene el formato esperado, se salta

        proveedor_nombre = NOMBRE_PROVEEDOR_SIN_IDENTIFICAR if nombre_hoja == "SIN IDENTIFICAR" else nombre_hoja
        proveedor_id = None  # se crea recién si hay al menos una fila para cargar

        for fila in filas_iter:
            if fila is None or all(c is None or str(c).strip() == "" for c in fila):
                continue

            def get(campo):
                i = idx.get(campo)
                return fila[i] if i is not None and i < len(fila) else None

            descripcion = _limpiar(get("descripcion"))
            if not descripcion:
                continue
            fila_categoria = _limpiar(get("categoria")) or categoria
            codigo_prov = _limpiar(get("codigo")) or None
            codigo_barras = _limpiar(get("codigo_barras"))
            costo = _numero(get("costo"))

            match, metodo, score = None, None, None
            if codigo_barras and codigo_barras in por_codigo_barras:
                match = next(c for c in candidatos if c["id"] == por_codigo_barras[codigo_barras])
                metodo = "codigo_barras"
            else:
                match, score = _matchear_por_descripcion(descripcion, fila_categoria, candidatos, umbral)
                if match:
                    metodo = "descripcion"

            if match:
                resumen["por_codigo_barras" if metodo == "codigo_barras" else "por_descripcion"] += 1
                if aplicar:
                    if proveedor_id is None:
                        proveedor_id = _obtener_o_crear_proveedor(conn, proveedor_nombre, resumen)
                    conn.execute(
                        """INSERT OR REPLACE INTO producto_proveedor (producto_id, proveedor_id, precio_costo, codigo_proveedor)
                           VALUES (?, ?, ?, ?)""",
                        (match["id"], proveedor_id, costo, codigo_prov),
                    )
                filas_reporte.append((proveedor_nombre, codigo_prov, descripcion, metodo, match["nombre"], round(score or 1.0, 2), costo))
            else:
                resumen["sin_match"] += 1
                filas_reporte.append((proveedor_nombre, codigo_prov, descripcion, "sin_match", "", round(score or 0, 2), costo))

        if aplicar:
            conn.commit()

    conn.close()

    _guardar_reporte(filas_reporte)

    print("=" * 60)
    print("  MATCHING DE PRODUCTOS ENTRE PROVEEDORES" + ("" if aplicar else "  (modo simulación, sin --aplicar)"))
    print("=" * 60)
    print(f"  Matcheados por código de barras: {resumen['por_codigo_barras']}")
    print(f"  Matcheados por descripción:      {resumen['por_descripcion']}")
    print(f"  Sin match (revisar a mano):      {resumen['sin_match']}")
    if resumen["proveedores_nuevos"]:
        print(f"  Proveedores nuevos creados: {', '.join(resumen['proveedores_nuevos'])}")
    print(f"  Reporte detallado: {REPORTE_POR_DEFECTO}")
    print("=" * 60)
    if not aplicar:
        print("\nEsto fue una simulación: no se escribió nada en producto_proveedor.")
        print("Revisá el reporte y, si el resultado se ve bien, volvé a correr con --aplicar.")


def _guardar_reporte(filas):
    os.makedirs(os.path.dirname(REPORTE_POR_DEFECTO), exist_ok=True)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Matching"
    ws.append(["Proveedor", "Código proveedor", "Descripción (proveedor)", "Método", "Producto matcheado", "Similitud", "Precio costo"])
    for fila in filas:
        ws.append(fila)
    wb.save(REPORTE_POR_DEFECTO)


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description="Puebla producto_proveedor matcheando listas de precios de proveedores contra el catálogo ya cargado.")
    parser.add_argument("archivo", help="Excel con una hoja por proveedor")
    parser.add_argument("--categoria", default=None, help="Limitar el piloto a una categoría (ej: 'Frenos')")
    parser.add_argument("--subcategoria", default=None, help="Limitar el piloto a una subcategoría (ej: 'Pastillas')")
    parser.add_argument("--umbral", type=float, default=UMBRAL_TEXTO_DEFAULT, help=f"Similitud mínima por descripción (default {UMBRAL_TEXTO_DEFAULT})")
    parser.add_argument("--aplicar", action="store_true", help="Sin esto corre en modo simulación (no escribe nada en la base)")
    args = parser.parse_args()

    procesar(args.archivo, args.categoria, args.subcategoria, args.umbral, args.aplicar)
